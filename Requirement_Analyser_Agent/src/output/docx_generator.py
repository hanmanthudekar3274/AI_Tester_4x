import io
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocxGenerator:
    """Converts a Markdown test plan string into a Word .docx file in memory."""

    def generate(self, markdown: str, title: str = "Test Plan") -> bytes:
        doc = Document()
        self._set_styles(doc)
        doc.add_heading(title, level=0)

        for line in markdown.splitlines():
            self._process_line(doc, line)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------

    def _process_line(self, doc: Document, line: str):
        stripped = line.rstrip()

        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("---"):
            doc.add_paragraph("─" * 60)
        elif re.match(r"^(\*|-|\d+\.)\s", stripped):
            text = re.sub(r"^(\*|-|\d+\.)\s", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            self._add_inline(p, text)
        elif "|" in stripped and stripped.startswith("|"):
            self._handle_table_row(doc, stripped)
        elif stripped == "":
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            self._add_inline(p, stripped)

    def _add_inline(self, paragraph, text: str):
        """Handle **bold** and `code` inline markdown."""
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            else:
                paragraph.add_run(part)

    _table_cache: list[list[str]] = []
    _current_table_rows: list[list[str]] = []

    def _handle_table_row(self, doc: Document, line: str):
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in cells):
            # Separator row — flush as Word table
            if self._current_table_rows:
                self._flush_table(doc)
        else:
            self._current_table_rows.append(cells)

    def _flush_table(self, doc: Document):
        if not self._current_table_rows:
            return
        headers = self._current_table_rows[0]
        rows = self._current_table_rows[1:]
        col_count = len(headers)
        table = doc.add_table(rows=1 + len(rows), cols=col_count)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            for run in hdr_cells[i].paragraphs[0].runs:
                run.bold = True
        for ri, row in enumerate(rows):
            row_cells = table.rows[ri + 1].cells
            for ci, cell_text in enumerate(row):
                if ci < col_count:
                    row_cells[ci].text = cell_text
        self._current_table_rows = []
        doc.add_paragraph()

    def _set_styles(self, doc: Document):
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
