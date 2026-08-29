"""Render a validated plan as a DOCX document.

Implements architecture/output_formatter_sop.md. Same content, same order as
format_markdown.py -- the two must not drift.

Returns bytes so the Streamlit download button can serve it without touching
the filesystem.
"""

from __future__ import annotations

import io
from typing import Any

DASH = "—"


def _text(value: Any) -> str:
    if value is None or value == "":
        return DASH
    return str(value).strip()


def _as_lines(items: Any) -> list[str]:
    if not isinstance(items, list):
        return []
    lines: list[str] = []
    for item in items:
        if isinstance(item, dict):
            name = item.get("name") or item.get("risk") or ""
            notes = item.get("notes") or item.get("mitigation") or ""
            lines.append(f"{_text(name)}: {notes}".strip(": ").strip())
        else:
            lines.append(_text(item))
    return [line for line in lines if line and line != DASH]


def render(plan: dict[str, Any], validation: dict[str, Any] | None = None,
           model: str | None = None) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError(
            "The python-docx package is not installed. Run: pip install python-docx"
        ) from exc

    plan = plan or {}
    validation = validation or {}
    stats = validation.get("stats", {})

    doc = Document()
    doc.add_heading(_text(plan.get("title")), level=0)

    meta = doc.add_paragraph()
    meta.add_run(f"Issue: {_text(plan.get('source_issue'))}\n").bold = True
    for label, key in (("Author", "author"), ("Generated", "generated_at"), ("Link", "url")):
        if plan.get(key):
            meta.add_run(f"{label}: {plan[key]}\n")
    if model:
        meta.add_run(f"Model: {model}\n")
    for run in meta.runs:
        run.font.size = Pt(9)

    # Coverage summary
    if stats:
        doc.add_heading("Coverage Summary", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text, header[1].text = "Metric", "Value"
        rows = [
            ("Test cases", stats.get("test_case_count", 0)),
            ("Acceptance criteria", stats.get("ac_count", 0)),
            ("Criteria covered", stats.get("ac_covered", 0)),
            ("Coverage", f"{stats.get('coverage_pct', 0)}%"),
            ("Inferred cases", stats.get("inferred_count", 0)),
        ]
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = str(label)
            cells[1].text = str(value)

    # Gaps above the test cases, per the SOP.
    gaps = _as_lines(plan.get("coverage_gaps"))
    if gaps:
        doc.add_heading("Coverage Gaps", level=1)
        doc.add_paragraph(
            "The issue did not contain enough information to test these points. "
            "They need a decision before this plan is complete."
        )
        for gap in gaps:
            doc.add_paragraph(gap, style="List Bullet")

    warnings = _as_lines(validation.get("warnings"))
    if warnings:
        doc.add_heading("Quality Warnings", level=1)
        for warning in warnings:
            doc.add_paragraph(warning, style="List Bullet")

    if plan.get("overview"):
        doc.add_heading("Overview", level=1)
        doc.add_paragraph(_text(plan["overview"]))

    scope = plan.get("scope")
    if not isinstance(scope, dict):
        scope = {}
    if scope.get("in_scope") or scope.get("out_of_scope"):
        doc.add_heading("Scope", level=1)
        for label, key in (("In scope", "in_scope"), ("Out of scope", "out_of_scope")):
            lines = _as_lines(scope.get(key))
            if lines:
                doc.add_paragraph(label, style="Intense Quote")
                for line in lines:
                    doc.add_paragraph(line, style="List Bullet")

    for heading, key in (
        ("Assumptions", "assumptions"),
        ("Test Environments", "test_environments"),
        ("Test Data Requirements", "test_data_requirements"),
        ("Entry Criteria", "entry_criteria"),
        ("Exit Criteria", "exit_criteria"),
    ):
        lines = _as_lines(plan.get(key))
        if lines:
            doc.add_heading(heading, level=1)
            for line in lines:
                doc.add_paragraph(line, style="List Bullet")

    if plan.get("test_strategy"):
        doc.add_heading("Test Strategy", level=1)
        doc.add_paragraph(_text(plan["test_strategy"]))

    _render_cases(doc, plan.get("test_cases"))
    _render_risks(doc, plan.get("risks"))
    _render_matrix(doc, plan.get("traceability_matrix"))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _render_cases(doc: Any, cases: Any) -> None:
    if not isinstance(cases, list) or not cases:
        return
    doc.add_heading("Test Cases", level=1)

    for case in cases:
        if not isinstance(case, dict):
            continue
        marker = "[INFERRED] " if case.get("inferred") else ""
        doc.add_heading(
            f"{_text(case.get('tc_id'))} {DASH} {marker}{_text(case.get('title'))}", level=2
        )
        doc.add_paragraph(
            f"Type: {_text(case.get('type'))}   |   "
            f"Priority: {_text(case.get('priority'))}   |   "
            f"Traces to: {_text(case.get('traces_to'))}"
        )
        if case.get("preconditions"):
            doc.add_paragraph(f"Preconditions: {_text(case['preconditions'])}")

        steps = case.get("steps")
        if isinstance(steps, list) and steps:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            header = table.rows[0].cells
            header[0].text, header[1].text, header[2].text = "#", "Action", "Test Data"
            for index, step in enumerate(steps, start=1):
                if not isinstance(step, dict):
                    continue
                cells = table.add_row().cells
                cells[0].text = str(step.get("step_no", index))
                cells[1].text = _text(step.get("action"))
                cells[2].text = _text(step.get("test_data"))
        else:
            doc.add_paragraph("No steps provided.")

        doc.add_paragraph(f"Expected Result: {_text(case.get('expected_result'))}")


def _render_risks(doc: Any, risks: Any) -> None:
    if not isinstance(risks, list) or not risks:
        return
    doc.add_heading("Risks", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, label in enumerate(("Risk", "Severity", "Mitigation", "Source")):
        header[index].text = label
    for risk in risks:
        cells = table.add_row().cells
        if isinstance(risk, dict):
            cells[0].text = _text(risk.get("risk"))
            cells[1].text = _text(risk.get("severity"))
            cells[2].text = _text(risk.get("mitigation"))
            cells[3].text = _text(risk.get("source"))
        else:
            cells[0].text = _text(risk)


def _render_matrix(doc: Any, matrix: Any) -> None:
    if not isinstance(matrix, list) or not matrix:
        return
    doc.add_heading("Traceability Matrix", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, label in enumerate(("AC", "Acceptance Criterion", "Test Cases")):
        header[index].text = label
    for row in matrix:
        if not isinstance(row, dict):
            continue
        cells = table.add_row().cells
        cells[0].text = _text(row.get("ac_id"))
        cells[1].text = _text(row.get("acceptance_criterion"))
        ids = row.get("test_case_ids")
        cells[2].text = ", ".join(str(i) for i in ids) if isinstance(ids, list) and ids else DASH
